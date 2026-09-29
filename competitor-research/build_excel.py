"""Bouwt Competitor_Research.xlsx uit data/competitors.csv en data/bestsellers.csv.

Gebruik:  python3 build_excel.py
Daarna:   python3 <xlsx-skill>/scripts/recalc.py Competitor_Research.xlsx
"""
import csv, datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).parent
OUT = BASE / "Competitor_Research.xlsx"
COUNTRY_NAMES = {"DE": "Duitsland", "FR": "Frankrijk", "UK": "Verenigd Koninkrijk",
                 "IT": "Italië", "ES": "Spanje", "NL": "Nederland", "AT": "Oostenrijk", "CH": "Zwitserland"}
COUNTRY_ORDER = ["DE", "FR", "UK", "IT", "ES", "NL", "AT", "CH"]

FONT = "Arial"
NAVY = "1F2A44"
f_base = Font(name=FONT, size=10)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_sub = Font(name=FONT, size=10, italic=True, color="666666")
f_bold = Font(name=FONT, size=10, bold=True)
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="1F5FBF", underline="single")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_input = PatternFill("solid", fgColor="FFF2CC")
fill_band = PatternFill("solid", fgColor="F5F7FA")
thin = Side(style="thin", color="D9DDE3")
border = Border(bottom=thin)

STATUS_COLORS = {"Prioriteit": ("C6EFCE", "006100"), "Monitoren": ("FFEB9C", "7F6000"),
                 "Onder drempel": ("F2F2F2", "595959"), "Geen SW-cijfer": ("F2F2F2", "595959"),
                 "Referentie": ("DDEBF7", "1F4E78")}

COLS = [  # (header, csv field, width, number format)
    ("Land", "Land", 7, None),
    ("Store", "Store", 24, None),
    ("Website", "Domein", 28, None),
    ("Categorie", "Categorie", 30, None),
    ("Bezoekers / mnd", "Bezoekers per maand", 15, '#,##0;-#,##0;"-"'),
    ("Status", None, 15, None),
    ("Trend vorige mnd", "Trend vorige maand", 14, '+0%;-0%;0%'),
    ("Match-score (1-5)", "Match-score", 11, "0"),
    ("Paid search", "Paid search", 11, "0%"),
    ("Top landen", "Top landen", 26, None),
    ("Netwerk / operator", "Netwerk / operator", 30, None),
    ("Gevonden via", "Gevonden via", 20, None),
    ("Eerst gevonden", "Eerst gevonden", 13, "yyyy-mm-dd"),
    ("Laatst gecheckt", "Laatst gecheckt", 13, "yyyy-mm-dd"),
    ("Producten gekopieerd", "Producten gekopieerd", 13, None),
    ("Notities", "Notities", 60, None),
]
COL = {h: get_column_letter(i + 1) for i, (h, *_) in enumerate(COLS)}
# Drempels staan op het Overzicht-tabblad (gele cellen) en worden overal gebruikt.
TH_PRIO, TH_MON, TH_SCORE = "Overzicht!$D$5", "Overzicht!$D$6", "Overzicht!$D$7"


def conv(field, val):
    if val in ("", None):
        return None
    if field in ("Bezoekers per maand", "Match-score"):
        return int(float(val))
    if field in ("Trend vorige maand", "Paid search"):
        return float(val)
    if field in ("Eerst gevonden", "Laatst gecheckt"):
        return datetime.date.fromisoformat(val)
    return val


def load():
    rows = list(csv.DictReader(open(BASE / "data/competitors.csv", encoding="utf-8")))
    key = lambda r: (COUNTRY_ORDER.index(r["Land"]) if r["Land"] in COUNTRY_ORDER else 99,
                     -(int(r["Bezoekers per maand"]) if r["Bezoekers per maand"] else -1), r["Store"])
    return sorted(rows, key=key)


def status_formula(r):
    v, s = f"{COL['Bezoekers / mnd']}{r}", f"{COL['Match-score (1-5)']}{r}"
    return (f'=IF(AND({s}<>"",{s}<{TH_SCORE}),"Referentie",IF({v}="","Geen SW-cijfer",'
            f'IF({v}>={TH_PRIO},"Prioriteit",IF({v}>={TH_MON},"Monitoren","Onder drempel"))))')


def header(ws, row, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill = f_head, fill_head
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 30


def competitor_sheet(wb, title, subtitle, rows):
    ws = wb.create_sheet(title)
    ws["A1"], ws["A1"].font = title, f_title
    ws["A2"], ws["A2"].font = subtitle, f_sub
    H = 4
    header(ws, H, [c[0] for c in COLS], [c[2] for c in COLS])
    ynn = DataValidation(type="list", formula1='"nee,ja,deels"', allow_blank=True)
    ws.add_data_validation(ynn)
    for i, row in enumerate(rows):
        r = H + 1 + i
        for j, (h, field, _, fmt) in enumerate(COLS, 1):
            c = ws.cell(row=r, column=j)
            c.value = status_formula(r) if h == "Status" else conv(field, row[field])
            c.font, c.border = f_base, border
            c.alignment = Alignment(vertical="top", wrap_text=(h in ("Notities", "Categorie", "Netwerk / operator")))
            if fmt:
                c.number_format = fmt
            if h == "Website" and c.value:
                c.hyperlink, c.font = f"https://{c.value}", f_link
            if h == "Producten gekopieerd":
                c.font, c.fill = f_input, fill_input
                ynn.add(c)
            if h == "Status":
                c.font = f_bold
    last = H + max(len(rows), 1)
    ws.auto_filter.ref = f"A{H}:{get_column_letter(len(COLS))}{last}"
    ws.freeze_panes = ws.cell(row=H + 1, column=3)
    st = COL["Status"]
    for name, (bg, fg) in STATUS_COLORS.items():
        ws.conditional_formatting.add(f"{st}{H+1}:{st}{last}", FormulaRule(
            formula=[f'${st}{H+1}="{name}"'], fill=PatternFill("solid", fgColor=bg), font=Font(name=FONT, bold=True, color=fg)))
    v = COL["Bezoekers / mnd"]
    ws.conditional_formatting.add(f"{v}{H+1}:{v}{last}", DataBarRule(start_type="num", start_value=0, end_type="max", color="8EA9DB"))
    tr = COL["Trend vorige mnd"]
    ws.conditional_formatting.add(f"{tr}{H+1}:{tr}{last}", FormulaRule(formula=[f"AND(${tr}{H+1}<>\"\",${tr}{H+1}>=0.25)"], font=Font(name=FONT, color="006100", bold=True)))
    ws.conditional_formatting.add(f"{tr}{H+1}:{tr}{last}", FormulaRule(formula=[f"AND(${tr}{H+1}<>\"\",${tr}{H+1}<=-0.25)"], font=Font(name=FONT, color="C00000")))
    ws.sheet_view.zoomScale = 100
    return ws


def overview(wb, countries, all_rows_last):
    ws = wb.active
    ws.title = "Overzicht"
    ws["A1"], ws["A1"].font = "Competitor Research – Google-dropshipping", f_title
    ws["A2"] = f"Laatst bijgewerkt: {datetime.date.today():%d-%m-%Y}  ·  bron bezoekers: SimilarWeb (schatting per maand)"
    ws["A2"].font = f_sub
    ws["A4"], ws["A4"].font = "Instellingen (gele cellen mag je aanpassen)", f_bold
    for r, (label, val, fmt) in enumerate([("Drempel prioriteit (bezoekers/mnd)", 75000, "#,##0"),
                                           ("Drempel monitoren (bezoekers/mnd)", 25000, "#,##0"),
                                           ("Minimale match-score (lager = referentie)", 3, "0")], 5):
        ws.cell(row=r, column=1, value=label).font = f_base
        c = ws.cell(row=r, column=4, value=val)
        c.font, c.fill, c.number_format = f_input, fill_input, fmt
    H = 10
    heads = ["Land", "Naam", "Stores totaal", "Prioriteit (≥ drempel)", "Monitoren", "Onder drempel / geen cijfer",
             "Referentie", "Bezoekers prioriteit (som)", "Grootste competitor (bezoekers)", "Nog niet gekopieerd (prioriteit)"]
    header(ws, H, heads, [8, 22, 13, 15, 12, 17, 12, 18, 18, 18])
    A = "'Alle competitors'!"
    rng = lambda h: f"{A}${COL[h]}$5:${COL[h]}${all_rows_last}"
    land, stat, vis, cop = rng("Land"), rng("Status"), rng("Bezoekers / mnd"), rng("Producten gekopieerd")
    for i, cc in enumerate(countries):
        r = H + 1 + i
        vals = [cc, COUNTRY_NAMES.get(cc, cc),
                f"=COUNTIFS({land},A{r})",
                f'=COUNTIFS({land},A{r},{stat},"Prioriteit")',
                f'=COUNTIFS({land},A{r},{stat},"Monitoren")',
                f'=COUNTIFS({land},A{r},{stat},"Onder drempel")+COUNTIFS({land},A{r},{stat},"Geen SW-cijfer")',
                f'=COUNTIFS({land},A{r},{stat},"Referentie")',
                f'=SUMIFS({vis},{land},A{r},{stat},"Prioriteit")',
                f'=_xlfn.MAXIFS({vis},{land},A{r},{stat},"<>Referentie")',
                f'=COUNTIFS({land},A{r},{stat},"Prioriteit",{cop},"nee")']
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border = f_base, border
            if j >= 8:
                c.number_format = "#,##0"
        ws.cell(row=r, column=1).hyperlink = f"#'{cc}'!A1"
        ws.cell(row=r, column=1).font = f_link
    tr = H + 1 + len(countries)
    ws.cell(row=tr, column=1, value="Totaal").font = f_bold
    for j in range(3, 11):
        colL = get_column_letter(j)
        f = f"=MAX({colL}{H+1}:{colL}{tr-1})" if j == 9 else f"=SUM({colL}{H+1}:{colL}{tr-1})"
        c = ws.cell(row=tr, column=j, value=f)
        c.font, c.number_format = f_bold, "#,##0"
        c.border = Border(top=Side(style="thin", color=NAVY))
    ws.cell(row=tr + 2, column=1, value="Klik op een landcode om naar dat tabblad te gaan.").font = f_sub
    ws.freeze_panes = "A11"


def bestsellers(wb):
    ws = wb.create_sheet("Bestsellers")
    ws["A1"], ws["A1"].font = "Bestsellers van prioriteit-competitors", f_title
    ws["A2"] = "Volgorde volgens de store zelf (sort_by=best-selling, over de hele looptijd, dus ook buiten het seizoen)."
    ws["A2"].font = f_sub
    heads = ["Land", "Store", "Rank", "Product", "Prijs", "Categorie", "URL", "Opgehaald op"]
    header(ws, 4, heads, [7, 24, 7, 60, 9, 28, 14, 13])
    rows = list(csv.DictReader(open(BASE / "data/bestsellers.csv", encoding="utf-8")))
    for i, b in enumerate(rows):
        r = 5 + i
        vals = [b["Land"], b["Domein"], int(b["Rank"]), b["Product"], float(b["Prijs"]), b["Categorie"], "Bekijk", datetime.date.fromisoformat(b["Opgehaald op"])]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border = f_base, border
        ws.cell(row=r, column=5).number_format = "0.00"
        ws.cell(row=r, column=8).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=7).hyperlink, ws.cell(row=r, column=7).font = b["URL"], f_link
    ws.auto_filter.ref = f"A4:H{4 + max(len(rows), 1)}"
    ws.freeze_panes = "A5"


def legend(wb):
    ws = wb.create_sheet("Legenda")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    ws["A1"], ws["A1"].font = "Legenda", f_title
    items = [
        ("Status", "Wordt automatisch berekend uit bezoekers en match-score, met de drempels op het Overzicht."),
        ("  Prioriteit", "Bezoekers ≥ drempel prioriteit: opnemen en bestsellers direct overnemen."),
        ("  Monitoren", "Tussen de twee drempels: opnemen en volgende maand opnieuw checken."),
        ("  Onder drempel", "Minder bezoekers dan de monitor-drempel."),
        ("  Geen SW-cijfer", "SimilarWeb toont geen bezoekersaantal (meestal < ~10K per maand)."),
        ("  Referentie", "Match-score onder het minimum: groot, maar ander model (bijv. eigen merk)."),
        ("Match-score", "1–5: hoe sterk de store op onze stores lijkt (Shopify, prijzen op ,95, lange SEO-titels, breed D&H-assortiment, veel paid search)."),
        ("Paid search", "Aandeel betaald zoekverkeer volgens SimilarWeb; hoog = vooral Google Shopping/Ads."),
        ("Trend", "Verandering in bezoekers t.o.v. de maand ervoor. Groen = ≥ +25%, rood = ≤ −25%."),
        ("Netwerk / operator", "Stores die vermoedelijk van dezelfde eigenaar zijn (zelfde thema, catalogus of naam)."),
        ("Producten gekopieerd", "Gele invulcel (nee / ja / deels): zelf bijhouden."),
        ("Bron", "SimilarWeb-websitepagina's (schatting), Tranco-ranglijst, Shopify /products.json van de stores zelf."),
        ("Bijwerken", "Vraag Claude: 'stuur de competitor sheet' of 'voeg nieuwe competitors toe voor <land>'."),
    ]
    for i, (k, v) in enumerate(items, 3):
        ws.cell(row=i, column=1, value=k).font = f_bold if not k.startswith("  ") else f_base
        ws.cell(row=i, column=2, value=v).font = f_base
        ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=5, column=1).fill = PatternFill("solid", fgColor="C6EFCE")
    ws.cell(row=6, column=1).fill = PatternFill("solid", fgColor="FFEB9C")
    ws.cell(row=9, column=1).fill = PatternFill("solid", fgColor="DDEBF7")


def main():
    rows = load()
    countries = [c for c in COUNTRY_ORDER if any(r["Land"] == c for r in rows)] + \
                sorted({r["Land"] for r in rows} - set(COUNTRY_ORDER))
    wb = Workbook()
    overview(wb, countries, 4 + len(rows))
    competitor_sheet(wb, "Alle competitors", f"{len(rows)} stores · gesorteerd op land en bezoekers", rows)
    for cc in countries:
        sub = [r for r in rows if r["Land"] == cc]
        competitor_sheet(wb, cc, f"{COUNTRY_NAMES.get(cc, cc)} · {len(sub)} stores", sub)
    bestsellers(wb)
    legend(wb)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = {"Overzicht": NAVY, "Alle competitors": "8EA9DB", "Bestsellers": "A9D08E", "Legenda": "BFBFBF"}.get(ws.title, "F4B183")
    wb.save(OUT)
    print(f"Opgeslagen: {OUT} ({len(rows)} competitors, landen: {', '.join(countries)})")


if __name__ == "__main__":
    main()
