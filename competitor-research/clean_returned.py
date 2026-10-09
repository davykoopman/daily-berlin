"""Leest een teruggestuurde Launch File, schoont statussen/listers/opmerkingen op en schrijft ze
in listings/launch_state.json. Kolommen worden op kopnaam gelezen (listers mogen kolommen toevoegen).

Regels (afgesproken met Davy):
  * Statussen: Not listed · Draft · Live · Issues · Duplicate   (Listed = Live)
  * Klaar = Live, Issues, Duplicate. Draft = voorbereid, nog niet live. Open = Not listed.
  * Status gezet maar geen lister → Dhafnie.
  * 'Already in/on (the) store' → status Duplicate, opmerking 'Already in store'.
  * 'Same product/item as #…' → Duplicate.
  * Not listed maar volgens fotovergelijking al in de store → Duplicate ('Already in store (auto-check)').

Gebruik:  python3 clean_returned.py <teruggestuurde.xlsx> [--autocheck]
"""
import json, re, sys
from pathlib import Path

BASE = Path(__file__).parent
STATE = BASE / "listings/launch_state.json"
STATUS_MAP = {"listed": "Live", "live": "Live", "not listed": "Not listed", "notlisted": "Not listed", "draft": "Draft",
              "issues": "Issues", "issue": "Issues", "duplicate": "Duplicate", "": "Not listed", "none": "Not listed"}
ALREADY = re.compile(r"already\s+(in|on)\s+(the\s+)?store", re.I)
SAME = re.compile(r"same\s+(item|product)", re.I)


def norm_status(v):
    return STATUS_MAP.get(str(v or "").strip().lower(), str(v).strip().capitalize())


def read_file(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        if not ws.title.endswith("– Links"):
            continue
        hdr = [str(c.value).strip() if c.value else "" for c in ws[3]]
        ix = {h: i for i, h in enumerate(hdr) if h}
        rows = {}
        for r in ws.iter_rows(min_row=4, values_only=True):
            link = r[ix["Link"]] if "Link" in ix else None
            if not link:
                continue
            g = lambda k: (r[ix[k]] if k in ix and ix[k] < len(r) else None)
            rows[str(link).strip()] = dict(name=g("Product Name"), lister=g("Lister"), status=g("Status"), comment=g("Comment"))
        out[ws.title.split(" – ")[0]] = rows
    return out


def clean(rec):
    st = norm_status(rec["status"])
    com = str(rec["comment"]).strip() if rec["comment"] not in (None, "") else ""
    if ALREADY.search(com):
        st, com = "Duplicate", "Already in store"
    elif SAME.search(com) and st in ("Issues", "Not listed", "Duplicate"):
        st = "Duplicate"
    lister = str(rec["lister"]).strip() if rec["lister"] else ""
    if st != "Not listed" and not lister:
        lister = "Dhafnie"
    if st == "Not listed" and not com:
        lister = lister if lister in ("Dhafnie", "Fatima", "Davy") else ""
    return dict(status=st, lister=lister, comment=com, name=str(rec["name"]).strip() if rec["name"] else "")


def main():
    path = sys.argv[1]
    state = json.loads(STATE.read_text())
    data = read_file(path)
    report = {}
    for store, rows in state.items():
        got = data.get(store[:22], data.get(store, {}))
        c = {}
        for x in rows:
            if x["url"] in got:
                x.update(clean(got[x["url"]]))
            c[x["status"]] = c.get(x["status"], 0) + 1
        report[store] = c
    if "--autocheck" in sys.argv:
        autocheck(state)
        for store, rows in state.items():
            report[store] = {}
            for x in rows:
                report[store][x["status"]] = report[store].get(x["status"], 0) + 1
    STATE.write_text(json.dumps(state, ensure_ascii=False))
    for s, c in report.items():
        print(s, c)


def autocheck(state):
    """Not listed-regels die al in de store staan (foto) → Duplicate."""
    import imagehash
    import product_picker as pp
    import select_products as sp
    cfg = json.loads((BASE / "config/stores.json").read_text())
    dom = {s["naam"]: s["domein"] for s in cfg["stores"]}
    hs = pp.hash_images([x["img"] for v in state.values() for x in v if x.get("img")])
    for store, rows in state.items():
        own = sp.own_catalog(dom[store])
        oh = [imagehash.hex_to_hash(h) for h in own["hashes"]]
        n = 0
        for x in rows:
            h = hs.get(x.get("img"))
            if x["status"] == "Not listed" and h and any(imagehash.hex_to_hash(h) - o <= 8 for o in oh):
                x.update(status="Duplicate", comment="Already in store (auto-check)", lister="")
                n += 1
        print(f"  autocheck {store}: {n} Not listed → Duplicate", file=sys.stderr)


if __name__ == "__main__":
    main()
