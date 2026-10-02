"""Product Picker: kiest bij een competitor de producten die het waard zijn om te listen.

Combineert per product:
  * hoe nieuw het is (publicatiedatum van de store zelf),
  * zijn plek in de bestseller-ranking van de store (sort_by=best-selling),
  * of de categorie past bij het seizoen van JOUW markt (config/seasons.json),
  * of je hem al hebt (fotovergelijking met je eigen store).

Gebruik:
  python3 product_picker.py karlson-berlin.com --markt DE --eigen dailyberlin.de
  python3 product_picker.py moreau-lyon.fr --markt ES --eigen jouwstore.es --dagen 45

Output: picks/<competitor>_<markt>_<datum>.xlsx + .txt met alleen de links.
"""
import argparse, concurrent.futures as cf, datetime as dt, io, json, math, re, ssl, sys, urllib.request
from pathlib import Path

BASE = Path(__file__).parent
CACHE = BASE / ".cache"
CTX = ssl.create_default_context(cafile="/root/.ccr/ca-bundle.crt") if Path("/root/.ccr/ca-bundle.crt").exists() else ssl.create_default_context()
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/128 Safari/537.36"}
# Volgorde is belangrijk: specifieker eerst (Strickjacke = vest, niet jas).
CAT_ORDER = ["Badmode", "Sandalen", "Shorts", "Vest / cardigan", "Winterjas", "Blazer", "Laarzen", "Tussenjas",
             "Trui / hoodie", "Set", "Jurk", "Broek", "Blouse / shirt", "Schoenen", "T-shirt / top"]


def get(url, binary=False, timeout=30):
    data = urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout, context=CTX).read()
    return data if binary else data.decode("utf8", "ignore")


def fetch_products(domain):
    out = []
    for page in range(1, 41):
        batch = json.loads(get(f"https://{domain}/products.json?limit=250&page={page}"))["products"]
        out += batch
        if len(batch) < 250:
            break
    return out


def fetch_bestseller_rank(domain, n_products):
    def page(n):
        try:
            h = get(f"https://{domain}/collections/all?sort_by=best-selling&page={n}")
        except Exception:
            return []
        seen = []
        for m in re.findall(r'href="(?:/[a-z]{2}(?:-[a-z]{2})?)?/(?:collections/[^/"]+/)?products/([^"?#/]+)', h):
            if m not in seen:
                seen.append(m)
        return seen
    first = page(1)
    per = max(len(first), 1)
    pages = [first]
    with cf.ThreadPoolExecutor(8) as ex:
        pages += list(ex.map(page, range(2, math.ceil(n_products / per) + 2)))
    rank = {}
    for p in pages:
        for h in p:
            rank.setdefault(h, len(rank) + 1)
    return rank


def sales_tiers(prods, rank):
    """Shopify sorteert producten met evenveel verkopen op datum (nieuwste eerst).
    Elke keer dat de publicatiedatum in de ranking weer 'nieuwer' wordt, begint een nieuwe
    groep met minder verkopen. Tier 0 = laatste groep = vrijwel zeker nog geen verkopen."""
    byh = {p["handle"]: p for p in prods}
    seq = [byh[h] for h, _ in sorted(rank.items(), key=lambda t: t[1]) if h in byh]
    if not seq:
        return {}
    groups = [[seq[0]]]
    for a, b in zip(seq, seq[1:]):
        (groups.append([b]) if b["published_at"] > a["published_at"] else groups[-1].append(b))
    return {p["handle"]: len(groups) - 1 - i for i, g in enumerate(groups) for p in g}


def compile_keywords(cfg):
    pats = {}
    for cat, kws in cfg["categories"].items():
        parts = [(r"\b" + re.escape(k.strip()) + r"\b") if len(k.strip()) <= 4 else re.escape(k) for k in kws]
        pats[cat] = re.compile("|".join(parts), re.I)
    return pats


def categorize(p, pats):
    # Titel eerst (Strickjacke = vest), dan product_type en tags.
    for text in (p["title"], p.get("product_type") or "", " ".join(p.get("tags") or [])):
        for cat in CAT_ORDER:
            if text and pats[cat].search(text):
                return cat
    return "Overig"


def season_weight(p, cat, cfg, markt, season):
    w_country = cfg["weights"].get(markt) or cfg["weights"][cfg["weights_fallback"].get(markt, "DE")]
    w = w_country[season].get(cat, 0.4)
    text = f'{p["title"]} {" ".join(p.get("tags") or [])}'.lower()
    tags = {cfg["competitor_season_tags"].get(t.lower()) for t in p.get("tags") or []} - {None}
    note = []
    if tags:  # de competitor labelt zelf het seizoen: sterk signaal
        if season in tags or (season == "autumn" and "winter" in tags) or (season == "winter" and "autumn" in tags):
            w = min(1.0, w + 0.2); note.append(f"tag {'/'.join(sorted(tags))}")
        elif tags <= {"summer"} and season in ("autumn", "winter"):
            w *= 0.3; note.append("tag summer")
    if any(k in text for k in cfg["penalty_keywords"]) and season in ("autumn", "winter"):
        w *= 0.5; note.append("zomer-woord")
    if any(k in text for k in cfg["boost_keywords"]) and season in ("autumn", "winter"):
        w = min(1.0, w + 0.1)
    return round(w, 2), ", ".join(note)


def img_url(p, width=240):
    if not p.get("images"):
        return None
    u = p["images"][0]["src"]
    return u + ("&" if "?" in u else "?") + f"width={width}"


def hash_images(urls):
    import imagehash
    from PIL import Image
    CACHE.mkdir(exist_ok=True)
    cf_path = CACHE / "phash.json"
    cache = json.loads(cf_path.read_text()) if cf_path.exists() else {}
    todo = [u for u in urls if u and u.split("?")[0] not in cache]

    def one(u):
        try:
            im = Image.open(io.BytesIO(get(u, binary=True, timeout=20))).convert("RGB")
            return u.split("?")[0], str(imagehash.phash(im))
        except Exception:
            return u.split("?")[0], None
    with cf.ThreadPoolExecutor(24) as ex:
        for k, v in ex.map(one, todo):
            cache[k] = v
    cf_path.write_text(json.dumps(cache))
    return {u: cache.get(u.split("?")[0]) for u in urls if u}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("competitor")
    ap.add_argument("--markt", default="DE", help="land van JOUW store (bepaalt het seizoen)")
    ap.add_argument("--eigen", help="domein van je eigen store, om dubbele producten te markeren")
    ap.add_argument("--dagen", type=int, default=60, help="hoe recent 'nieuw' is (dagen)")
    ap.add_argument("--maand", type=int, default=dt.date.today().month)
    ap.add_argument("--min-seizoen", type=float, default=0.6)
    ap.add_argument("--max", type=int, default=120, help="max aantal picks")
    a = ap.parse_args()

    cfg = json.loads((BASE / "config/seasons.json").read_text())
    season = cfg["season_by_month"].get(a.markt, cfg["season_by_month"]["DE"])[str(a.maand)]
    pats = compile_keywords(cfg)
    today = dt.datetime.now(dt.timezone.utc)

    print(f"Producten ophalen van {a.competitor} ...", file=sys.stderr)
    prods = fetch_products(a.competitor)
    rank = fetch_bestseller_rank(a.competitor, len(prods))
    n = max(len(rank), 1)
    tiers = sales_tiers(prods, rank)

    rows = []
    for p in prods:
        live = max(dt.datetime.fromisoformat(p["published_at"] or p["created_at"]), dt.datetime.fromisoformat(p["created_at"]))
        age = max((today - live).days, 0)
        r = rank.get(p["handle"])
        pct = 1 - (r - 1) / n if r else 0.0          # 1.0 = beste verkoper
        cat = categorize(p, pats)
        sw, snote = season_weight(p, cat, cfg, a.markt, season)
        is_new = age <= a.dagen
        newness = max(0.0, 1 - age / max(a.dagen * 1.5, 1))
        tier = tiers.get(p["handle"])
        # Alleen de laatste groep is betrouwbaar als 'nog geen verkopen'; verder telt de ranking zelf.
        signal = ("nog geen verkopen" if tier == 0 else "top-verkoper" if pct >= 0.85 else
                  "verkoopt" if pct >= 0.6 else "weinig verkopen" if r else "onbekend")
        # Nieuw én al hoog in de ranking = snelle verkoper.
        hot = is_new and pct >= 0.75
        sales_bonus = -0.2 if signal == "nog geen verkopen" else 0.0
        score = sw * (0.55 * pct + 0.45 * newness + sales_bonus) + (0.15 if hot else 0)
        avail = [v for v in p["variants"] if v.get("available")]
        label = ("🔥 Nieuwe winner" if hot else "Nieuw" if is_new else
                 "Seizoens-bestseller" if pct >= 0.85 else "")
        rows.append(dict(p=p, cat=cat, season_w=sw, season_note=snote, age=age, live=live.date(), rank=r, pct=pct, signal=signal,
                         score=round(score, 3), label=label, price=float(p["variants"][0]["price"]),
                         compare=float(p["variants"][0]["compare_at_price"] or 0), sizes_av=f"{len(avail)}/{len(p['variants'])}",
                         url=f"https://{a.competitor}/products/{p['handle']}", img=img_url(p)))

    picks = [x for x in rows if x["label"] and x["season_w"] >= a.min_seizoen]
    picks.sort(key=lambda x: -x["score"])
    picks = picks[: a.max]
    print(f"{len(prods)} producten, {len(picks)} picks (seizoen {season}, markt {a.markt})", file=sys.stderr)

    dupes = {}
    if a.eigen:
        print(f"Vergelijken met {a.eigen} ...", file=sys.stderr)
        own = fetch_products(a.eigen)
        own_imgs = {img_url(p): p for p in own if p.get("images")}
        hashes = hash_images(list(own_imgs) + [x["img"] for x in picks])
        import imagehash
        own_h = [(imagehash.hex_to_hash(hashes[u]), p) for u, p in own_imgs.items() if hashes.get(u)]
        for x in picks:
            h = hashes.get(x["img"])
            if not h:
                continue
            h = imagehash.hex_to_hash(h)
            best = min(own_h, key=lambda t: t[0] - h, default=None)
            if best and best[0] - h <= 10:
                dupes[x["url"]] = f'https://{a.eigen}/products/{best[1]["handle"]}'
    for x in picks:
        x["dupe"] = dupes.get(x["url"], "")

    out_dir = BASE / "picks"
    out_dir.mkdir(exist_ok=True)
    stem = f'{a.competitor.split(".")[0]}_{a.markt}_{dt.date.today():%Y-%m-%d}'
    write_excel(out_dir / f"{stem}.xlsx", picks, a, season, len(prods))
    (out_dir / f"{stem}.txt").write_text("\n".join(x["url"] for x in picks if not x["dupe"]) + "\n")
    print(out_dir / f"{stem}.xlsx")


def write_excel(path, picks, a, season, n_total):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from PIL import Image

    F = "Arial"
    navy = "1F2A44"
    wb = Workbook()
    ws = wb.active
    ws.title = "Picks"
    seas_nl = {"autumn": "herfst", "winter": "winter", "spring": "lente", "summer": "zomer"}[season]
    ws["A1"] = f"Product picks: {a.competitor} → markt {a.markt} ({seas_nl})"
    ws["A1"].font = Font(name=F, size=15, bold=True, color=navy)
    ws["A2"] = (f"{len(picks)} van {n_total} producten geselecteerd · nieuw = live ≤ {a.dagen} dagen · "
                f"seizoensfit ≥ {a.min_seizoen} · gemaakt {dt.date.today():%d-%m-%Y}")
    ws["A2"].font = Font(name=F, size=10, italic=True, color="666666")
    heads = ["Foto", "Product", "Label", "Categorie", "Seizoensfit", "Live sinds", "Dagen live", "Bestseller-rank",
             "Top %", "Verkoopsignaal", "Prijs", "Was-prijs", "Maten op voorraad", "Score", "Al in eigen store?", "Listen?", "Link", "Opmerking"]
    widths = [14, 46, 18, 16, 11, 12, 10, 13, 8, 17, 9, 9, 12, 8, 22, 10, 12, 24]
    H = 4
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=H, column=i, value=h)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=navy)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    dv = DataValidation(type="list", formula1='"ja,nee"', allow_blank=True)
    ws.add_data_validation(dv)
    fills = {"🔥 Nieuwe winner": "FCE4D6", "Nieuw": "E2EFDA", "Seizoens-bestseller": "DDEBF7"}
    thin = Border(bottom=Side(style="thin", color="D9DDE3"))

    def thumb(u):
        try:
            im = Image.open(io.BytesIO(get(u, binary=True, timeout=20))).convert("RGB")
            im.thumbnail((90, 90))
            b = io.BytesIO(); im.save(b, "JPEG", quality=80); b.seek(0)
            return b
        except Exception:
            return None
    with cf.ThreadPoolExecutor(16) as ex:
        thumbs = list(ex.map(lambda x: thumb(x["img"]) if x["img"] else None, picks))

    for i, (x, tb) in enumerate(zip(picks, thumbs)):
        r = H + 1 + i
        ws.row_dimensions[r].height = 70
        vals = ["", x["p"]["title"], x["label"], x["cat"], x["season_w"], x["live"], x["age"], x["rank"], x["pct"], x["signal"],
                x["price"], x["compare"] or None, x["sizes_av"], x["score"],
                "ja" if x["dupe"] else "", "nee" if x["dupe"] else "", "Open", x["season_note"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name=F, size=10)
            c.alignment = Alignment(vertical="center", wrap_text=(j in (2, 18)))
            c.border = thin
        ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=fills.get(x["label"], "FFFFFF"))
        ws.cell(row=r, column=6).number_format = "dd-mm-yyyy"
        ws.cell(row=r, column=9).number_format = "0%"
        ws.cell(row=r, column=5).number_format = "0.0"
        for col in (11, 12):
            ws.cell(row=r, column=col).number_format = "0.00"
        lst = ws.cell(row=r, column=16)
        lst.fill, lst.font = PatternFill("solid", fgColor="FFF2CC"), Font(name=F, size=10, color="0000FF")
        dv.add(lst)
        link = ws.cell(row=r, column=17)
        link.hyperlink, link.font = x["url"], Font(name=F, size=10, color="1F5FBF", underline="single")
        if x["dupe"]:
            d = ws.cell(row=r, column=15)
            d.hyperlink, d.font = x["dupe"], Font(name=F, size=10, color="C00000", underline="single")
        if tb:
            img = XLImage(tb); img.anchor = f"A{r}"
            ws.add_image(img)
    ws.freeze_panes = ws.cell(row=H + 1, column=3)
    ws.auto_filter.ref = f"A{H}:R{H + max(len(picks), 1)}"

    ls = wb.create_sheet("Links")
    ls["A1"] = "Links van alle picks die je nog niet hebt (kopieer kolom A in je import-app)"
    ls["A1"].font = Font(name=F, bold=True)
    for i, x in enumerate([x for x in picks if not x["dupe"]], 3):
        ls.cell(row=i, column=1, value=x["url"]).font = Font(name=F, size=10)
        ls.cell(row=i, column=2, value=x["label"]).font = Font(name=F, size=10)
        ls.cell(row=i, column=3, value=x["cat"]).font = Font(name=F, size=10)
    ls.column_dimensions["A"].width = 80
    ls.column_dimensions["B"].width = 20
    ls.column_dimensions["C"].width = 18

    lg = wb.create_sheet("Uitleg")
    lg.column_dimensions["A"].width = 22
    lg.column_dimensions["B"].width = 100
    rows = [("Label", ""),
            ("  🔥 Nieuwe winner", f"Live ≤ {a.dagen} dagen én al in de top 25% bestsellers van de store: verkoopt snel."),
            ("  Nieuw", f"Live ≤ {a.dagen} dagen; nog niet bewezen, maar de competitor heeft er net op ingezet."),
            ("  Seizoens-bestseller", "Ouder product, maar in de top 15% bestsellers en past bij het seizoen."),
            ("Seizoensfit", "0–1: hoe goed de categorie past bij het seizoen van jouw markt (config/seasons.json). Tags van de competitor zelf (Winter/Summer) tellen mee."),
            ("Bestseller-rank", "Plek in 'sorteer op best verkocht' van de store zelf (1 = best verkocht)."),
            ("Verkoopsignaal", "Afgeleid uit de bestseller-ranking: producten met evenveel verkopen staan op datum gesorteerd. 'nog geen verkopen' = in de laatste groep (vrijwel zeker 0 verkopen); de andere labels volgen uit de positie in de ranking. Schatting, geen aantallen."),
            ("Maten op voorraad", "Aantal varianten dat nog leverbaar is. Veel uitverkocht kan juist duiden op goede verkoop."),
            ("Al in eigen store?", "Gevonden via fotovergelijking met je eigen store. 'ja' = link naar jouw product."),
            ("Listen?", "Gele cel: zet op 'ja' voor wat je wilt listen en filter daarop; kopieer dan de links.")]
    for i, (k, v) in enumerate(rows, 1):
        lg.cell(row=i, column=1, value=k).font = Font(name=F, size=10, bold=not k.startswith("  "))
        lg.cell(row=i, column=2, value=v).font = Font(name=F, size=10)
    wb.save(path)


if __name__ == "__main__":
    main()
